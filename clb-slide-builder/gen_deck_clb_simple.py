#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Pipeline đơn giản cho CLB Năng lực số và Ứng dụng AI
Đọc JSON → Tạo PPTX trực tiếp bằng python-pptx → Upload Drive
Không qua ppt-master, không có cổng kiểm tra phức tạp.
"""
import json
import sys
import hashlib
import subprocess
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Đường dẫn
GENERATED = Path(__file__).parent.parent
CLB_DECK = GENERATED / "clb_deck"
DRIVE_DIR = Path.home() / "Library/CloudStorage/GoogleDrive-mxuanvan159@gmail.com/My Drive/00_THPT_THUAN_HOA/12_CLB_NLS_AI/02_SLIDE"
LOGO_PATH = CLB_DECK / "logo_truong_alpha.png"


def add_title_slide(prs, d):
    """Slide tiêu đề"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d["ten_bai"]
    p.font.size = Pt(44)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(1))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"CLB Năng lực số và Ứng dụng AI\nTiết {d['tiet']} · Tuần {d['tuan']}"
    p.font.size = Pt(24)
    p.alignment = PP_ALIGN.CENTER
    
    # Footer
    footer_box = slide.shapes.add_textbox(Inches(0.5), Inches(6.5), Inches(9), Inches(0.5))
    tf = footer_box.text_frame
    p = tf.paragraphs[0]
    p.text = "CLB Năng lực số · GV Mai Xuân Văn"
    p.font.size = Pt(12)
    p.font.color.rgb = RGBColor(128, 128, 128)
    p.alignment = PP_ALIGN.CENTER


def add_khoi_dong_slide(prs, d):
    """Slide khởi động"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "KHỞI ĐỘNG"
    p.font.size = Pt(36)
    p.font.bold = True
    
    # Nội dung khởi động
    kd = d.get("khoi_dong", {})
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(4))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = kd.get("tieu_de", "")
    p.font.size = Pt(24)
    p.font.bold = True
    
    for the in kd.get("the", []):
        p = tf.add_paragraph()
        p.text = f"\n{the['title']}"
        p.font.size = Pt(20)
        p.font.bold = True
        for line in the.get("lines", []):
            p = tf.add_paragraph()
            p.text = f"  {line}"
            p.font.size = Pt(18)
    
    # Câu hỏi
    if kd.get("cau_hoi"):
        p = tf.add_paragraph()
        p.text = f"\n💬 {kd['cau_hoi']}"
        p.font.size = Pt(20)
        p.font.italic = True


def add_muc_tieu_slide(prs, d):
    """Slide mục tiêu"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "MỤC TIÊU BÀI HỌC"
    p.font.size = Pt(36)
    p.font.bold = True
    
    # Nội dung
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for mt in d.get("muc_tieu", []):
        p = tf.add_paragraph()
        p.text = f"✓ {mt['title']}"
        p.font.size = Pt(22)
        p.font.bold = True
        for line in mt.get("lines", []):
            p = tf.add_paragraph()
            p.text = f"    {line}"
            p.font.size = Pt(18)


def add_noi_dung_slide(prs, p_data):
    """Slide nội dung (1 slide cho mỗi phần)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = p_data.get("tieu_de", "").upper()
    p.font.size = Pt(32)
    p.font.bold = True
    
    # Sub
    if p_data.get("sub"):
        sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.3), Inches(9), Inches(0.5))
        tf = sub_box.text_frame
        p = tf.paragraphs[0]
        p.text = p_data["sub"]
        p.font.size = Pt(18)
        p.font.italic = True
    
    # Ghi bài
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(4.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for gb in p_data.get("ghi_bai", []):
        p = tf.add_paragraph()
        p.text = gb
        p.font.size = Pt(18)
    
    # Thẻ
    y_pos = 2 + len(p_data.get("ghi_bai", [])) * 0.4
    if y_pos < 4.5:
        y_pos = 4.5
    
    for the in p_data.get("the", []):
        card_box = slide.shapes.add_textbox(Inches(0.5), Inches(y_pos), Inches(9), Inches(1.5))
        tf = card_box.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"【{the['title']}】"
        p.font.size = Pt(18)
        p.font.bold = True
        
        for line in the.get("lines", []):
            p = tf.add_paragraph()
            p.text = f"  {line}"
            p.font.size = Pt(16)
        
        y_pos += 1.5


def add_luyen_tap_slide(prs, d):
    """Slide luyện tập"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "LUYỆN TẬP"
    p.font.size = Pt(36)
    p.font.bold = True
    
    # Nội dung
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for lv in d.get("luyentap", []):
        p = tf.add_paragraph()
        p.text = f"\n{lv['tieu_de']}"
        p.font.size = Pt(20)
        p.font.bold = True
        
        if lv.get("sub"):
            p = tf.add_paragraph()
            p.text = f"  {lv['sub']}"
            p.font.size = Pt(16)
            p.font.italic = True
        
        for viec in lv.get("viec", []):
            p = tf.add_paragraph()
            p.text = f"  {viec}"
            p.font.size = Pt(16)


def add_van_dung_slide(prs, d):
    """Slide vận dụng"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "VẬN DỤNG"
    p.font.size = Pt(36)
    p.font.bold = True
    
    # Nội dung
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for vd in d.get("vandung", []):
        p = tf.add_paragraph()
        p.text = f"\n{vd.get('nhan', '')}"
        p.font.size = Pt(20)
        p.font.bold = True
        
        for line in vd.get("lines", []):
            p = tf.add_paragraph()
            p.text = f"  {line}"
            p.font.size = Pt(18)


def add_tong_ket_slide(prs, d):
    """Slide tổng kết"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tiêu đề
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "TỔNG KẾT"
    p.font.size = Pt(36)
    p.font.bold = True
    
    # Nội dung
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9), Inches(5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = d.get("tong_ket", "")
    p.font.size = Pt(20)
    
    if d.get("loi_can_tranh"):
        p = tf.add_paragraph()
        p.text = "\n⚠ Lưu ý:"
        p.font.size = Pt(20)
        p.font.bold = True
        
        for loi in d.get("loi_can_tranh", []):
            p = tf.add_paragraph()
            p.text = f"  {loi}"
            p.font.size = Pt(18)


def add_ket_thuc_slide(prs, d):
    """Slide kết thúc"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Tóm tắt
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(1), Inches(9), Inches(1))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = "TÓM TẮT"
    p.font.size = Pt(32)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = d.get("tom_tat_vo", "")
    p.font.size = Pt(20)
    p.alignment = PP_ALIGN.CENTER
    
    # Về nhà
    home_box = slide.shapes.add_textbox(Inches(0.5), Inches(4), Inches(9), Inches(1.5))
    tf = home_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "🏠 VỀ NHÀ"
    p.font.size = Pt(28)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    p = tf.add_paragraph()
    p.text = d.get("ket_thuc", "")
    p.font.size = Pt(18)
    p.alignment = PP_ALIGN.CENTER
    
    # Dẫn dò
    next_box = slide.shapes.add_textbox(Inches(0.5), Inches(6), Inches(9), Inches(0.5))
    tf = next_box.text_frame
    p = tf.paragraphs[0]
    p.text = d.get("dan_do", "")
    p.font.size = Pt(14)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER


def main():
    if len(sys.argv) < 2:
        print("Cách dùng: gen_deck_clb_simple.py <noi_dung.json>")
        sys.exit(1)
    
    json_path = Path(sys.argv[1])
    if not json_path.is_absolute():
        json_path = CLB_DECK / json_path
    
    d = json.loads(json_path.read_text(encoding="utf-8"))
    bai = d["bai_so"]
    ten_bai = d["ten_bai"]
    ten_file = d["ten_file"]
    
    print(f"\n{'='*60}")
    print(f"PIPELINE CLB Bài {bai}: {ten_bai}")
    print(f"{'='*60}\n")
    
    # Tạo presentation
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Thêm các slide
    add_title_slide(prs, d)
    add_khoi_dong_slide(prs, d)
    add_muc_tieu_slide(prs, d)
    
    for p in d.get("phan", []):
        add_noi_dung_slide(prs, p)
    
    add_luyen_tap_slide(prs, d)
    add_van_dung_slide(prs, d)
    add_tong_ket_slide(prs, d)
    add_ket_thuc_slide(prs, d)
    
    # Lưu PPTX
    pptx_path = GENERATED / f"{ten_file}.pptx"
    prs.save(pptx_path)
    print(f"✓ Đã tạo: {pptx_path.name}")
    print(f"  Số slide: {len(prs.slides)}")
    
    # Upload Drive
    DRIVE_DIR.mkdir(parents=True, exist_ok=True)
    dest = DRIVE_DIR / f"{ten_file}.pptx"
    subprocess.run(["cp", str(pptx_path), str(dest)])
    
    h1 = hashlib.sha256(pptx_path.read_bytes()).hexdigest()
    h2 = hashlib.sha256(dest.read_bytes()).hexdigest()
    ok = h1 == h2
    
    print(f"✓ Upload: {dest.name} — SHA {'KHỚP ✓' if ok else 'LỆCH ✗'}")
    
    return ok


if __name__ == "__main__":
    success = main()
    sys.exit(0 if success else 1)
