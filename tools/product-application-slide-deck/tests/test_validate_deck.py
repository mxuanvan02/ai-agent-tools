#!/usr/bin/env python3
"""Test suite for product application slide deck validation."""

import sys
import unittest
import tempfile
import shutil
from pathlib import Path

# Add parent directory to path for imports
sys.path.insert(0, str(Path(__file__).parent.parent))

try:
    from pptx import Presentation
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

from scripts.validate_deck import (
    validate_slide_count,
    validate_bounds,
    validate_aspect_ratios,
    validate_speaker_notes
)


@unittest.skipUnless(HAS_PPTX, "python-pptx not installed")
class TestValidateDeck(unittest.TestCase):
    """Test deck validation functions."""

    def setUp(self):
        """Create temporary directory for test files."""
        self.temp_dir = tempfile.mkdtemp()
        self.test_pptx = Path(self.temp_dir) / "test.pptx"

    def tearDown(self):
        """Clean up temporary files."""
        shutil.rmtree(self.temp_dir)

    def _create_test_pptx(self, num_slides=15, with_notes=True, with_out_of_bounds=False):
        """Helper to create test PPTX files."""
        p = Presentation()
        p.slide_width = Inches(13.333)
        p.slide_height = Inches(7.5)

        for i in range(num_slides):
            slide = p.slides.add_slide(p.slide_layouts[6])

            # Add a text box
            tx = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
            tx.text_frame.text = f"Slide {i+1}"

            # Add speaker notes if requested
            if with_notes:
                slide.notes_slide.notes_text_frame.text = f"Notes for slide {i+1}"

            # Add out-of-bounds element if requested
            if with_out_of_bounds and i == 0:
                shape = slide.shapes.add_shape(
                    MSO_SHAPE.RECTANGLE,
                    Inches(14),  # Beyond right edge
                    Inches(1),
                    Inches(1),
                    Inches(1)
                )

        p.save(str(self.test_pptx))
        return self.test_pptx

    def test_validate_slide_count_pass(self):
        """Test slide count validation passes with correct count."""
        pptx_path = self._create_test_pptx(num_slides=15)
        self.assertTrue(validate_slide_count(pptx_path, expected=15))

    def test_validate_slide_count_fail(self):
        """Test slide count validation fails with wrong count."""
        pptx_path = self._create_test_pptx(num_slides=10)
        self.assertFalse(validate_slide_count(pptx_path, expected=15))

    def test_validate_bounds_pass(self):
        """Test bounds validation passes with no violations."""
        pptx_path = self._create_test_pptx(with_out_of_bounds=False)
        self.assertTrue(validate_bounds(pptx_path))

    def test_validate_bounds_fail(self):
        """Test bounds validation fails with out-of-bounds elements."""
        pptx_path = self._create_test_pptx(with_out_of_bounds=True)
        self.assertFalse(validate_bounds(pptx_path))

    def test_validate_speaker_notes_pass(self):
        """Test speaker notes validation passes when all slides have notes."""
        pptx_path = self._create_test_pptx(with_notes=True)
        self.assertTrue(validate_speaker_notes(pptx_path))

    def test_validate_speaker_notes_fail(self):
        """Test speaker notes validation fails when slides lack notes."""
        pptx_path = self._create_test_pptx(with_notes=False)
        self.assertFalse(validate_speaker_notes(pptx_path))


class TestPublicHygiene(unittest.TestCase):
    """Test public hygiene checks."""

    def test_imports(self):
        """Test that all required modules can be imported."""
        try:
            import scripts.validate_deck
            import scripts.create_product_deck
        except ImportError as e:
            self.fail(f"Failed to import module: {e}")


if __name__ == "__main__":
    unittest.main()
